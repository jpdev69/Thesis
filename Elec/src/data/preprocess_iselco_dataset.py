"""Preprocess ISUE-ISELCO raw workbook into thesis-ready CSV deliverables.

Outputs:
1) clean_account_month.csv
2) campus_month_aggregate.csv
3) feature_table_monthly.csv
4) data_quality_report.md
"""

from __future__ import annotations

import argparse
import math
import re
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ACCOUNT_PATTERN = re.compile(r"^\d{2}-\d{4}-\d{4}$")
MONTH_MAP = {
    "JAN": 1,
    "FEB": 2,
    "MAR": 3,
    "APR": 4,
    "MAY": 5,
    "JUN": 6,
    "JUL": 7,
    "AUG": 8,
    "SEP": 9,
    "OCT": 10,
    "NOV": 11,
    "DEC": 12,
}


@dataclass
class PipelineArtifacts:
    clean_account_month: pd.DataFrame
    campus_month_aggregate: pd.DataFrame
    feature_table_monthly: pd.DataFrame
    stats: dict


def _normalize_text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if text == "":
        return None
    return re.sub(r"\s+", " ", text)


def _to_number(value) -> float:
    if value is None:
        return np.nan
    if isinstance(value, (int, float, np.number)):
        return float(value)
    text = _normalize_text(value)
    if text is None:
        return np.nan
    text = text.replace(",", "")
    text = text.replace("PHP", "")
    text = text.replace("P", "")
    text = text.strip()
    try:
        return float(text)
    except ValueError:
        return np.nan


def _parse_month(value) -> float:
    if value is None:
        return np.nan
    if isinstance(value, (int, float, np.number)) and not math.isnan(float(value)):
        month_val = int(float(value))
        if 1 <= month_val <= 12:
            return float(month_val)
    text = _normalize_text(value)
    if text is None:
        return np.nan
    key = text.upper()[:3]
    month_val = MONTH_MAP.get(key)
    return float(month_val) if month_val is not None else np.nan


def _extract_raw_records(xlsx_path: Path, sheet_name: str = "Database") -> pd.DataFrame:
    wb = load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in workbook")

    ws = wb[sheet_name]
    records = []

    for row_idx in range(1, ws.max_row + 1):
        row_values = [ws.cell(row=row_idx, column=col_idx).value for col_idx in range(1, ws.max_column + 1)]
        row_values = [_normalize_text(v) if isinstance(v, str) else v for v in row_values]

        account_idx = None
        for i, value in enumerate(row_values):
            if isinstance(value, str) and ACCOUNT_PATTERN.match(value):
                account_idx = i
                break

        if account_idx is None:
            continue

        building = row_values[account_idx - 1] if account_idx - 1 >= 0 else None
        account_no = row_values[account_idx]
        year_raw = row_values[account_idx + 1] if account_idx + 1 < len(row_values) else None
        month_raw = row_values[account_idx + 2] if account_idx + 2 < len(row_values) else None
        kwh_raw = row_values[account_idx + 3] if account_idx + 3 < len(row_values) else None
        gross_raw = row_values[account_idx + 4] if account_idx + 4 < len(row_values) else None

        records.append(
            {
                "source_row": row_idx,
                "building_name": building,
                "account_no": account_no,
                "year_raw": year_raw,
                "month_raw": month_raw,
                "kwh_raw": kwh_raw,
                "gross_raw": gross_raw,
            }
        )

    if not records:
        raise ValueError("No records extracted from workbook")

    df = pd.DataFrame(records)
    return df


def _normalize_and_type(df_raw: pd.DataFrame) -> pd.DataFrame:
    df = df_raw.copy()

    df["building_name"] = df["building_name"].apply(_normalize_text)
    df["building_name"] = df["building_name"].fillna("UNKNOWN")

    df["year"] = pd.to_numeric(df["year_raw"], errors="coerce")
    df["year"] = df["year"].round().astype("Int64")

    df["month"] = df["month_raw"].apply(_parse_month).astype("Int64")

    df["kwh_raw"] = df["kwh_raw"].apply(_to_number)
    df["gross_raw"] = df["gross_raw"].apply(_to_number)

    valid_period_mask = df["year"].notna() & df["month"].notna()
    df["period"] = pd.NaT
    df.loc[valid_period_mask, "period"] = pd.to_datetime(
        {
            "year": df.loc[valid_period_mask, "year"].astype(int),
            "month": df.loc[valid_period_mask, "month"].astype(int),
            "day": 1,
        },
        errors="coerce",
    )

    df["has_kwh_raw"] = df["kwh_raw"].notna()
    df["has_gross_raw"] = df["gross_raw"].notna()
    df["raw_completeness_score"] = df[["has_kwh_raw", "has_gross_raw"]].sum(axis=1)

    # Guardrails for invalid values
    df.loc[df["kwh_raw"] <= 0, "kwh_raw"] = np.nan
    df.loc[df["gross_raw"] < 0, "gross_raw"] = np.nan

    return df


def _deduplicate_account_period(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    with_period = df[df["period"].notna()].copy()
    without_period = df[df["period"].isna()].copy()

    with_period = with_period.sort_values(
        by=["account_no", "period", "raw_completeness_score", "source_row"],
        ascending=[True, True, False, False],
    )

    group_sizes = with_period.groupby(["account_no", "period"]).size().rename("dedup_group_size")
    with_period = with_period.merge(group_sizes, on=["account_no", "period"], how="left")

    with_period["dedup_rank"] = (
        with_period.groupby(["account_no", "period"]).cumcount() + 1
    )

    kept = with_period[with_period["dedup_rank"] == 1].copy()
    dropped = with_period[with_period["dedup_rank"] > 1].copy()

    kept["had_duplicate_in_group"] = kept["dedup_group_size"] > 1
    kept["is_duplicate_dropped"] = False
    dropped["had_duplicate_in_group"] = True
    dropped["is_duplicate_dropped"] = True

    if not without_period.empty:
        without_period = without_period.copy()
        without_period["dedup_group_size"] = 1
        without_period["dedup_rank"] = 1
        without_period["had_duplicate_in_group"] = False
        without_period["is_duplicate_dropped"] = False
        kept = pd.concat([kept, without_period], ignore_index=True, sort=False)

    kept = kept.sort_values(["account_no", "period", "source_row"]).reset_index(drop=True)
    dropped = dropped.sort_values(["account_no", "period", "source_row"]).reset_index(drop=True)
    return kept, dropped


def _clean_and_impute(df: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    clean = df.copy()

    clean["unit_cost_raw"] = np.where(
        clean["kwh_raw"].notna() & clean["gross_raw"].notna() & (clean["kwh_raw"] > 0),
        clean["gross_raw"] / clean["kwh_raw"],
        np.nan,
    )

    period_unit_cost = clean.groupby("period")["unit_cost_raw"].median()
    global_unit_cost = float(clean["unit_cost_raw"].median(skipna=True))
    clean["unit_cost_reference"] = clean["period"].map(period_unit_cost)
    clean["unit_cost_reference"] = clean["unit_cost_reference"].fillna(global_unit_cost)

    clean["kwh_clean"] = clean["kwh_raw"]
    clean["gross_clean"] = clean["gross_raw"]
    clean["is_imputed_kwh"] = False
    clean["is_imputed_gross"] = False

    can_impute_gross = (
        clean["kwh_clean"].notna()
        & clean["gross_clean"].isna()
        & clean["unit_cost_reference"].notna()
        & (clean["kwh_clean"] > 0)
    )
    clean.loc[can_impute_gross, "gross_clean"] = (
        clean.loc[can_impute_gross, "kwh_clean"] * clean.loc[can_impute_gross, "unit_cost_reference"]
    )
    clean.loc[can_impute_gross, "is_imputed_gross"] = True

    can_impute_kwh = (
        clean["gross_clean"].notna()
        & clean["kwh_clean"].isna()
        & clean["unit_cost_reference"].notna()
        & (clean["unit_cost_reference"] > 0)
    )
    clean.loc[can_impute_kwh, "kwh_clean"] = (
        clean.loc[can_impute_kwh, "gross_clean"] / clean.loc[can_impute_kwh, "unit_cost_reference"]
    )
    clean.loc[can_impute_kwh, "is_imputed_kwh"] = True

    clean["has_kwh_clean"] = clean["kwh_clean"].notna()
    clean["has_gross_clean"] = clean["gross_clean"].notna()
    clean["has_both_clean"] = clean["has_kwh_clean"] & clean["has_gross_clean"]

    clean["unit_cost_clean"] = np.where(
        clean["has_both_clean"] & (clean["kwh_clean"] > 0),
        clean["gross_clean"] / clean["kwh_clean"],
        np.nan,
    )

    valid_unit_cost = clean["unit_cost_clean"].dropna()
    q1 = float(valid_unit_cost.quantile(0.25))
    q3 = float(valid_unit_cost.quantile(0.75))
    iqr = q3 - q1
    lower = max(0.0, q1 - 1.5 * iqr)
    upper = q3 + 1.5 * iqr

    clean["is_unit_cost_outlier"] = (
        clean["unit_cost_clean"].notna()
        & ((clean["unit_cost_clean"] < lower) | (clean["unit_cost_clean"] > upper))
    )

    clean["unit_cost_capped"] = clean["unit_cost_clean"].clip(lower=lower, upper=upper)
    clean["gross_capped"] = np.where(
        clean["kwh_clean"].notna() & clean["unit_cost_capped"].notna(),
        clean["kwh_clean"] * clean["unit_cost_capped"],
        np.nan,
    )

    clean["record_status"] = "missing_both"
    clean.loc[clean["has_kwh_raw"] & clean["has_gross_raw"], "record_status"] = "complete_raw"
    clean.loc[clean["is_imputed_gross"], "record_status"] = "imputed_gross"
    clean.loc[clean["is_imputed_kwh"], "record_status"] = "imputed_kwh"

    clean = clean.sort_values(["account_no", "period", "source_row"]).reset_index(drop=True)

    stats = {
        "global_unit_cost_median": global_unit_cost,
        "unit_cost_outlier_lower": lower,
        "unit_cost_outlier_upper": upper,
        "imputed_kwh_count": int(clean["is_imputed_kwh"].sum()),
        "imputed_gross_count": int(clean["is_imputed_gross"].sum()),
        "rows_with_both_clean": int(clean["has_both_clean"].sum()),
        "unit_cost_outlier_count": int(clean["is_unit_cost_outlier"].sum()),
    }
    return clean, stats


def _build_campus_month_aggregate(clean: pd.DataFrame) -> pd.DataFrame:
    usable = clean[clean["period"].notna()].copy()
    if usable.empty:
        raise ValueError("No usable period rows found after cleaning")

    period_min = usable["period"].min()
    period_max = usable["period"].max()
    full_periods = pd.period_range(period_min, period_max, freq="M").to_timestamp()

    total_accounts_reference = int(usable["account_no"].nunique())

    monthly = (
        usable.groupby("period", as_index=False)
        .agg(
            active_meter_count=("has_both_clean", "sum"),
            complete_raw_meter_count=("raw_completeness_score", lambda x: int((x == 2).sum())),
            imputed_meter_count=("record_status", lambda x: int(x.isin(["imputed_kwh", "imputed_gross"]).sum())),
            outlier_meter_count=("is_unit_cost_outlier", "sum"),
            total_kwh_raw_complete=("kwh_raw", lambda x: float(np.nansum(x))),
            total_gross_raw_complete=("gross_raw", lambda x: float(np.nansum(x))),
            total_kwh_clean=("kwh_clean", lambda x: float(np.nansum(x))),
            total_gross_clean=("gross_clean", lambda x: float(np.nansum(x))),
            total_gross_capped=("gross_capped", lambda x: float(np.nansum(x))),
        )
    )

    monthly = monthly.set_index("period").reindex(full_periods).reset_index().rename(columns={"index": "period"})

    count_cols = [
        "active_meter_count",
        "complete_raw_meter_count",
        "imputed_meter_count",
        "outlier_meter_count",
    ]
    for col in count_cols:
        monthly[col] = monthly[col].fillna(0).astype(int)

    sum_cols = [
        "total_kwh_raw_complete",
        "total_gross_raw_complete",
        "total_kwh_clean",
        "total_gross_clean",
        "total_gross_capped",
    ]
    for col in sum_cols:
        monthly[col] = monthly[col].replace(0, np.nan)

    monthly["total_accounts_reference"] = total_accounts_reference
    monthly["coverage_ratio"] = monthly["active_meter_count"] / total_accounts_reference

    monthly["avg_unit_cost_clean"] = monthly["total_gross_clean"] / monthly["total_kwh_clean"]
    monthly["avg_unit_cost_capped"] = monthly["total_gross_capped"] / monthly["total_kwh_clean"]
    monthly["kwh_per_active_meter"] = monthly["total_kwh_clean"] / monthly["active_meter_count"].replace(0, np.nan)
    monthly["gross_per_active_meter"] = monthly["total_gross_clean"] / monthly["active_meter_count"].replace(0, np.nan)
    monthly["gross_capped_per_active_meter"] = monthly["total_gross_capped"] / monthly["active_meter_count"].replace(0, np.nan)

    monthly["year"] = monthly["period"].dt.year
    monthly["month"] = monthly["period"].dt.month

    cols = [
        "period",
        "year",
        "month",
        "total_accounts_reference",
        "active_meter_count",
        "coverage_ratio",
        "complete_raw_meter_count",
        "imputed_meter_count",
        "outlier_meter_count",
        "total_kwh_raw_complete",
        "total_gross_raw_complete",
        "total_kwh_clean",
        "total_gross_clean",
        "total_gross_capped",
        "avg_unit_cost_clean",
        "avg_unit_cost_capped",
        "kwh_per_active_meter",
        "gross_per_active_meter",
        "gross_capped_per_active_meter",
    ]
    return monthly[cols].sort_values("period").reset_index(drop=True)


def _add_lag_features(df: pd.DataFrame, col: str, prefix: str, lags: list[int]) -> pd.DataFrame:
    for lag in lags:
        df[f"{prefix}_lag_{lag}"] = df[col].shift(lag)
    return df


def _build_feature_table(campus: pd.DataFrame) -> pd.DataFrame:
    feat = campus.copy().sort_values("period").reset_index(drop=True)

    feat["month_num"] = feat["period"].dt.month
    feat["quarter"] = feat["period"].dt.quarter
    feat["month_sin"] = np.sin(2 * np.pi * feat["month_num"] / 12)
    feat["month_cos"] = np.cos(2 * np.pi * feat["month_num"] / 12)

    feat = _add_lag_features(feat, "total_kwh_clean", "kwh", [1, 2, 3, 6, 12])
    feat = _add_lag_features(feat, "total_gross_capped", "gross", [1, 2, 3, 6, 12])
    feat = _add_lag_features(feat, "avg_unit_cost_capped", "unit_cost", [1, 3, 6, 12])

    feat["kwh_mom_pct"] = feat["total_kwh_clean"].pct_change()
    feat["gross_mom_pct"] = feat["total_gross_capped"].pct_change()
    feat["unit_cost_mom_pct"] = feat["avg_unit_cost_capped"].pct_change()

    feat["kwh_roll_mean_3"] = feat["total_kwh_clean"].shift(1).rolling(window=3, min_periods=3).mean()
    feat["kwh_roll_mean_6"] = feat["total_kwh_clean"].shift(1).rolling(window=6, min_periods=6).mean()
    feat["kwh_roll_std_3"] = feat["total_kwh_clean"].shift(1).rolling(window=3, min_periods=3).std()
    feat["kwh_roll_std_6"] = feat["total_kwh_clean"].shift(1).rolling(window=6, min_periods=6).std()

    feat["gross_roll_mean_3"] = feat["total_gross_capped"].shift(1).rolling(window=3, min_periods=3).mean()
    feat["gross_roll_mean_6"] = feat["total_gross_capped"].shift(1).rolling(window=6, min_periods=6).mean()
    feat["unit_cost_roll_mean_3"] = feat["avg_unit_cost_capped"].shift(1).rolling(window=3, min_periods=3).mean()

    feat = feat.replace([np.inf, -np.inf], np.nan)

    feat["target_kwh_next_month"] = feat["total_kwh_clean"].shift(-1)
    feat["target_gross_next_month"] = feat["total_gross_capped"].shift(-1)
    feat["target_unit_cost_next_month"] = feat["avg_unit_cost_capped"].shift(-1)

    required = [
        "kwh_lag_1",
        "kwh_lag_2",
        "kwh_lag_3",
        "kwh_roll_mean_3",
        "gross_lag_1",
        "unit_cost_lag_1",
        "target_kwh_next_month",
    ]
    feat["is_trainable_row"] = feat[required].notna().all(axis=1)

    return feat


def _render_report(stats: dict, output_dir: Path) -> str:
    clean_path = output_dir / "clean_account_month.csv"
    campus_path = output_dir / "campus_month_aggregate.csv"
    feature_path = output_dir / "feature_table_monthly.csv"

    return f"""# Data Quality Report and Rationale

## Scope

This report documents preprocessing applied to `ISUE_ISELCO_Monitoring.xlsx` to produce thesis-ready model inputs.

## Deliverable 1: `clean_account_month.csv`

### What it contains

- One cleaned row per `(account_no, period)` after deduplication.
- Raw and cleaned values for `kwh` and `gross`.
- Imputation flags, outlier flags, and capping fields.
- Record-level status labels (`complete_raw`, `imputed_kwh`, `imputed_gross`, `missing_both`).

### Why these steps were taken

- **Account-pattern extraction** avoids accidental inclusion of title/header noise from the workbook.
- **Canonical typing** (`year`, `month`, `period`, numeric values) ensures deterministic time-series ordering and reproducibility.
- **Deduplication by `(account, period)`** prevents double-counting in campus totals.
- **Conservative imputation** (using period/global median unit cost) fills only one-sided missing pairs while preserving traceability via flags.
- **Unit-cost outlier capping** protects downstream modeling from extreme billing anomalies while retaining original values for audit.

### Run statistics

- Extracted rows: {stats['rows_extracted']}
- Rows with valid period: {stats['rows_with_valid_period']}
- Unique accounts: {stats['unique_accounts']}
- Duplicate rows removed: {stats['duplicate_rows_removed']}
- Kept rows after dedup: {stats['rows_after_dedup']}
- Imputed `kwh` rows: {stats['imputed_kwh_count']}
- Imputed `gross` rows: {stats['imputed_gross_count']}
- Rows with both cleaned values: {stats['rows_with_both_clean']}
- Unit-cost outlier rows: {stats['unit_cost_outlier_count']}
- Unit-cost IQR cap range: [{stats['unit_cost_outlier_lower']:.4f}, {stats['unit_cost_outlier_upper']:.4f}]

## Deliverable 2: `campus_month_aggregate.csv`

### What it contains

- Monthly campus totals (`total_kwh_clean`, `total_gross_clean`, `total_gross_capped`).
- Coverage controls (`active_meter_count`, `coverage_ratio`, reference account count).
- Data quality diagnostics per month (raw-complete meter count, imputed meter count, outlier meter count).
- Intensity metrics (`kwh_per_active_meter`, `gross_per_active_meter`).

### Why these steps were taken

- **Monthly aggregation** aligns raw billing granularity with practical forecasting horizon available from this dataset.
- **Coverage features** separate true demand movement from reporting-coverage movement.
- **Parallel raw/clean/capped totals** support defendable comparisons between strict raw evidence and model-stable features.

### Run statistics

- Monthly periods covered: {stats['aggregate_periods']}
- First period: {stats['aggregate_start']}
- Last period: {stats['aggregate_end']}
- Periods with active observations: {stats['aggregate_periods_with_activity']}

## Deliverable 3: `feature_table_monthly.csv`

### What it contains

- Model-ready engineered features from cleaned monthly series.
- Seasonality encodings (`month_sin`, `month_cos`), lag features, MoM change, and rolling statistics.
- Next-month targets for `kwh`, `gross`, and `unit_cost`.
- `is_trainable_row` indicator for rows meeting minimum feature completeness.

### Why these steps were taken

- **Lag features** capture autocorrelation and recurring billing behavior.
- **Rolling statistics** capture local trend/volatility better than single lags alone.
- **Seasonality encoding** enables models to learn annual cyclic behavior without ordinal-month distortion.
- **Separate next-month targets** supports direct forecasting setups and two-stage cost forecasting.

### Run statistics

- Feature rows: {stats['feature_rows']}
- Trainable rows: {stats['feature_trainable_rows']}

## Deliverable 4: `data_quality_report.md`

### What it contains

- End-to-end rationale, assumptions, and measurable outcomes for each transformation stage.

### Why this is necessary for thesis defense

- It provides **method transparency** (what changed and why).
- It preserves **auditability** (raw vs cleaned paths and flags).
- It supports **defensibility** of preprocessing choices during panel questioning.

## Output Locations

- `{clean_path}`
- `{campus_path}`
- `{feature_path}`
"""


def run_pipeline(input_path: Path, output_dir: Path) -> PipelineArtifacts:
    raw = _extract_raw_records(input_path, sheet_name="Database")
    normalized = _normalize_and_type(raw)
    deduped, dropped = _deduplicate_account_period(normalized)
    cleaned, clean_stats = _clean_and_impute(deduped)
    campus = _build_campus_month_aggregate(cleaned)
    features = _build_feature_table(campus)

    stats = {
        "rows_extracted": int(len(normalized)),
        "rows_with_valid_period": int(normalized["period"].notna().sum()),
        "unique_accounts": int(normalized["account_no"].nunique()),
        "duplicate_rows_removed": int(len(dropped)),
        "rows_after_dedup": int(len(deduped)),
        "imputed_kwh_count": clean_stats["imputed_kwh_count"],
        "imputed_gross_count": clean_stats["imputed_gross_count"],
        "rows_with_both_clean": clean_stats["rows_with_both_clean"],
        "unit_cost_outlier_count": clean_stats["unit_cost_outlier_count"],
        "unit_cost_outlier_lower": clean_stats["unit_cost_outlier_lower"],
        "unit_cost_outlier_upper": clean_stats["unit_cost_outlier_upper"],
        "aggregate_periods": int(len(campus)),
        "aggregate_start": campus["period"].min().strftime("%Y-%m"),
        "aggregate_end": campus["period"].max().strftime("%Y-%m"),
        "aggregate_periods_with_activity": int((campus["active_meter_count"] > 0).sum()),
        "feature_rows": int(len(features)),
        "feature_trainable_rows": int(features["is_trainable_row"].sum()),
    }

    return PipelineArtifacts(
        clean_account_month=cleaned,
        campus_month_aggregate=campus,
        feature_table_monthly=features,
        stats=stats,
    )


def main() -> None:
    base_dir = Path(__file__).resolve().parents[3]
    default_input = base_dir / "ISUE_ISELCO_Monitoring.xlsx"
    default_output = base_dir / "Elec" / "data" / "processed"
    default_report = base_dir / "Elec" / "docs" / "data_quality_report.md"

    parser = argparse.ArgumentParser(description="Preprocess ISUE-ISELCO billing dataset")
    parser.add_argument("--input", type=Path, default=default_input, help="Path to raw xlsx file")
    parser.add_argument("--output-dir", type=Path, default=default_output, help="Directory for CSV outputs")
    parser.add_argument("--report-path", type=Path, default=default_report, help="Path for markdown report")
    args = parser.parse_args()

    artifacts = run_pipeline(args.input, args.output_dir)

    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.report_path.parent.mkdir(parents=True, exist_ok=True)

    clean_path = args.output_dir / "clean_account_month.csv"
    campus_path = args.output_dir / "campus_month_aggregate.csv"
    feature_path = args.output_dir / "feature_table_monthly.csv"

    artifacts.clean_account_month.to_csv(clean_path, index=False)
    artifacts.campus_month_aggregate.to_csv(campus_path, index=False)
    artifacts.feature_table_monthly.to_csv(feature_path, index=False)

    report_text = _render_report(artifacts.stats, args.output_dir)
    args.report_path.write_text(report_text, encoding="utf-8")

    print("Preprocessing completed.")
    print(f"- {clean_path}")
    print(f"- {campus_path}")
    print(f"- {feature_path}")
    print(f"- {args.report_path}")


if __name__ == "__main__":
    main()
